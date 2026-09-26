#!/usr/bin/env python3
"""md2docx: one Bearing Markdown artifact as a client-ready Word document.

The document opens with a cover block (title, project, customer, prepared
by, version, date), the version history table and "About this document",
then the Markdown body: headings, paragraphs with bold, italic and code,
bullet and numbered lists, pipe tables as real tables, fenced code in a
shaded monospace block, and images. Template guidance comments (<!-- -->)
are dropped. A Mermaid block is never pasted as text: when a rendered image
is named for it (an image line right after the block, or --diagram for the
first diagram only) the
image goes in; otherwise a one-line pointer to the source file does. An
SVG image is swapped for the PNG of the same name, since Word cannot show
SVG.

Needs python-docx: run it as
  uv run --quiet --with python-docx python3 md2docx.py ...

Usage: md2docx.py <source.md> <out.docx> --title T --project P
                  [--customer C] [--prepared-by E] [--version N]
                  [--date YYYY-MM-DD] [--about TEXT]
                  [--history history.json] [--diagram path.png]
Prints a counts line; exits 1 on a missing or empty source.
"""

import argparse
import datetime
import json
import os
import re
import sys

try:
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm
except ImportError:  # pack.py reports this instead of failing the pack
    docx = None

INK = "1F2937"
ACCENT = "2F4A7A"
SHADE = "EEF2F7"


def shade(cell_or_par, colour):
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), colour)
    pr.append(shd)


def runs(par, text, size=None, colour=None, bold=False):
    """Add text with **bold**, *italic*, `code` and [links](url) as runs."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    for part in re.split(
        r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|_[^_\s][^_]*_)", text
    ):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif (part.startswith("*") and part.endswith("*")) or (
            part.startswith("_") and part.endswith("_") and len(part) > 2
        ):
            r = par.add_run(part[1:-1])
            r.italic = True
        else:
            r = par.add_run(part)
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = RGBColor.from_string(colour)


def table(d, rows, header=True):
    cols = max(len(r) for r in rows)
    t = d.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, r in enumerate(rows):
        cells = t.add_row().cells
        for j in range(cols):
            p = cells[j].paragraphs[0]
            runs(p, r[j] if j < len(r) else "", size=9, bold=header and i == 0)
            if header and i == 0:
                shade(cells[j], SHADE)
    d.add_paragraph()
    return t


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def resolve_image(src_dir, path):
    """The image beside the source, else from the repository root; an SVG
    becomes the PNG of the same name."""
    for base in (src_dir, os.getcwd()):
        p = os.path.normpath(os.path.join(base, path))
        if p.lower().endswith(".svg"):
            p = p[:-4] + ".png"
        if os.path.exists(p):
            return p
    return None


def body(d, md, src_dir, diagram):
    """Render the Markdown body. Returns counts."""
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    lines = md.split("\n")
    counts = {
        "headings": 0,
        "tables": 0,
        "images": 0,
        "diagrams_as_pointer": 0,
        "code": 0,
    }
    i, first_h1 = 0, True
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            lang = s[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            if lang == "mermaid":
                nxt = next((ln for ln in lines[i : i + 3] if ln.strip()), "")
                m = re.match(r"^\s*!\[[^\]]*\]\(([^)]+)\)", nxt)
                img = (
                    resolve_image(src_dir, m.group(1))
                    if m
                    else (
                        diagram
                        if diagram and os.path.exists(diagram) and not counts["images"]
                        else None
                    )
                )
                if img:
                    d.add_picture(img, width=Cm(16))
                    counts["images"] += 1
                    if m:
                        i = lines.index(nxt, i) + 1
                else:
                    p = d.add_paragraph()
                    runs(
                        p,
                        "Diagram: drawn in the source document (Mermaid); open the Markdown file or the rendered diagram beside it.",
                        size=9,
                        colour="4B5563",
                    )
                    counts["diagrams_as_pointer"] += 1
                continue
            p = d.add_paragraph()
            shade(p, SHADE)
            r = p.add_run("\n".join(block))
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            counts["code"] += 1
            continue
        h = re.match(r"^(#{1,4})\s+(.*)$", s)
        if h:
            level = len(h.group(1))
            if level == 1 and first_h1:
                first_h1 = False  # the cover carries the title
                i += 1
                continue
            d.add_heading(re.sub(r"[`*]", "", h.group(2)).strip(), level=min(level, 3) if level > 1 else 1)
            counts["headings"] += 1
            i += 1
            continue
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s) or re.match(
            r"^(Diagram): `?([^\s`]+\.(?:svg|png))`?", s
        )
        if img:
            p = resolve_image(src_dir, img.group(2))
            if p:
                d.add_picture(p, width=Cm(16))
                cap = d.add_paragraph()
                runs(cap, img.group(1), size=9, colour="4B5563")
                counts["images"] += 1
            else:
                par = d.add_paragraph()
                runs(
                    par,
                    f"Image not found beside the source: {img.group(2)}",
                    size=9,
                    colour="B45309",
                )
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                table(d, rows)
                counts["tables"] += 1
            continue
        lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if lm:
            depth = len(lm.group(1).replace("\t", "  ")) // 2
            numbered = lm.group(2)[0].isdigit()
            style = ("List Number" if numbered else "List Bullet") + (
                f" {min(depth + 1, 3)}" if depth else ""
            )
            try:
                p = d.add_paragraph(style=style)
            except KeyError:
                p = d.add_paragraph(style="List Number" if numbered else "List Bullet")
            text = lm.group(3)
            i += 1
            while (
                i < len(lines)
                and lines[i].startswith("  ")
                and lines[i].strip()
                and not re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[i])
            ):
                text += " " + lines[i].strip()
                i += 1
            runs(p, text)
            continue
        if s.startswith(">"):
            p = d.add_paragraph()
            runs(p, s.lstrip("> "), colour="4B5563")
            p.paragraph_format.left_indent = Cm(0.8)
            i += 1
            continue
        para = [s]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not re.match(r"^\s*(#|\||```|!\[|>|[-*+]\s|\d+[.)]\s)", lines[i])
        ):
            para.append(lines[i].strip())
            i += 1
        p = d.add_paragraph()
        runs(p, " ".join(para))
    return counts


def footer(d, text):
    sec = d.sections[0]
    p = sec.footer.paragraphs[0]
    runs(p, text + "   ·   page ", size=8, colour="6B7280")
    r = p.add_run()
    for kind, val in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = val
            r._r.append(it)
    r.font.size = Pt(8)


def build(
    src,
    out,
    title,
    project,
    customer,
    prepared_by,
    version,
    date,
    about,
    history,
    diagram=None,
):
    if docx is None:
        raise RuntimeError(
            "python-docx is not installed (run with: uv run --with python-docx)"
        )
    md = open(src, encoding="utf-8").read()
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string(INK)
    for s in ("Heading 1", "Heading 2", "Heading 3"):
        d.styles[s].font.color.rgb = RGBColor.from_string(ACCENT)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)

    p = d.add_paragraph()
    runs(p, title, size=26, colour=ACCENT, bold=True)
    p = d.add_paragraph()
    runs(p, project, size=14, colour="4B5563")
    table(
        d,
        [
            ["Project", project],
            ["Customer", customer or "not recorded"],
            ["Prepared by", prepared_by or "not recorded"],
            ["Version", f"v{version}"],
            ["Date", date],
        ],
        header=False,
    )
    d.add_heading("Version history", level=2)
    rows = [["Version", "Date", "Author", "Summary of changes"]]
    for h in history or [
        {
            "version": version,
            "date": date,
            "author": prepared_by,
            "summary": "First issue",
        }
    ]:
        rows.append(
            [
                f"v{h['version']}",
                h.get("date", ""),
                h.get("author", ""),
                h.get("summary", ""),
            ]
        )
    table(d, rows)
    d.add_heading("About this document", level=2)
    p = d.add_paragraph()
    runs(
        p,
        about
        or f"The {title.lower()} for {project}, generated from {os.path.basename(src)} in the project repository.",
    )
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.add_page_break()
    counts = body(d, md, os.path.dirname(os.path.abspath(src)), diagram)
    footer(d, f"{project}  ·  {title}  v{version}")
    core = d.core_properties
    core.title, core.author, core.subject = (
        f"{project} {title} v{version}",
        prepared_by or "",
        title,
    )
    d.save(out)
    return counts


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("out")
    ap.add_argument("--title", required=True)
    ap.add_argument("--project", required=True)
    ap.add_argument("--customer", default="")
    ap.add_argument("--prepared-by", default="")
    ap.add_argument("--version", type=int, default=1)
    ap.add_argument("--date", default=datetime.date.today().isoformat())
    ap.add_argument("--about", default="")
    ap.add_argument("--history", default="")
    ap.add_argument("--diagram", default="")
    a = ap.parse_args()
    if (
        not os.path.isfile(a.source)
        or not open(a.source, encoding="utf-8").read().strip()
    ):
        print(
            f"md2docx: {a.source} is missing or empty, nothing written", file=sys.stderr
        )
        return 1
    history = json.load(open(a.history, encoding="utf-8")) if a.history else None
    try:
        c = build(
            a.source,
            a.out,
            a.title,
            a.project,
            a.customer,
            a.prepared_by,
            a.version,
            a.date,
            a.about,
            history,
            a.diagram or None,
        )
    except RuntimeError as e:
        print(f"md2docx: {e}", file=sys.stderr)
        return 1
    print(
        f"md2docx: {a.out}: {c['headings']} headings, {c['tables']} tables, {c['images']} images, "
        f"{c['code']} code blocks, {c['diagrams_as_pointer']} diagrams without an image"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
