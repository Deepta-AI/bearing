#!/usr/bin/env python3
"""pack: the client documentation pack, rebuilt in full from the repository.

Reads a manifest (docs/deliverables/pack.json, else the kit's default in
templates/pack.json) that says which Bearing artifact goes in which of the
twelve numbered folders, and writes deliverables/<Project>_ProjectDocumentation/:

  - each artifact as its Markdown source (HTML comments removed), plus <Project>_<Key>_v<N>.docx with
    the cover block and version history (md2docx.py), plus any CSV the
    manifest asks for (a Markdown table whose first header cell matches),
    plus copied companions (JSON models, SQL, diagrams, HTML screens);
  - a README per folder (what it holds, which skill fills it, its files);
  - the top README (what is where, the current version of every artifact)
    and CHANGELOG.md (every version, newest first);
  - <Project>_DeliveryChecklist_<date>.xlsx from templates/delivery-checklist.json;
  - a root-anchored "/deliverables/" rule in .gitignore (a bare "deliverables/"
    would also ignore docs/deliverables/, so it is rewritten).

Versions live in docs/deliverables/versions.json, which is committed: an
artifact whose sources changed gets the next version and a history row that
counts what changed (sections and ids added, expanded, retired), or the
--summary text. Unchanged sources keep their version, so rebuilding twice
bumps nothing.

The checklist keeps what people entered: docs/deliverables/checklist.json
holds each task's status, owner, target date and remarks. --import-checklist
<xlsx> reads a filled copy back in. A task is marked Completed from evidence
(a file its evidence glob finds, optionally containing a text after "|")
only when no person has set its status; a person's status is never changed.

Needs python-docx and openpyxl: run it as
  uv run --quiet --with python-docx --with openpyxl python3 pack.py ...

Usage: pack.py [--manifest M] [--out DIR] [--project P] [--customer C]
               [--prepared-by E] [--summary TEXT] [--date YYYY-MM-DD]
               [--import-checklist FILE.xlsx] [--no-docx]
Prints one line per artifact and a counts line; exits 1 when no artifact
source exists, when a Word document or the workbook could not be written,
or when the manifest names a folder outside the twelve.
"""

import argparse
import datetime
import glob
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
KIT_TEMPLATES = os.path.join(os.path.dirname(HERE), "templates")
sys.path.insert(0, HERE)
import md2docx  # noqa: E402

FOLDERS = [
    "01_Product",
    "02_Technical",
    "03_Architecture",
    "04_API_Documentation",
    "05_UI_UX",
    "06_QA",
    "07_Analytics",
    "08_Security",
    "09_Deployment",
    "10_Release_Notes",
    "11_Client_Deliverables",
    "12_Meeting_Notes",
]
ID_RE = re.compile(r"^(?:[A-Z]{1,6}-)+[A-Z0-9]+(?:-[A-Z0-9]+)*$")
LIST_ID_RE = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+\**([A-Z][A-Z0-9-]*[A-Z0-9])\**[:.\s]")


def read_json(path, default):
    try:
        return json.load(open(path, encoding="utf-8"))
    except (OSError, ValueError):
        return default


def pascal(s):
    return "".join(w[:1].upper() + w[1:] for w in re.split(r"[^A-Za-z0-9]+", s) if w)


def expand(patterns):
    out = []
    for p in patterns if isinstance(patterns, list) else [patterns]:
        out += sorted(f for f in glob.glob(p, recursive=True) if os.path.isfile(f))
    return out


def fingerprint(paths):
    """Hash of the sources, their sections, and a hash per id: the first cell
    of a table row, or the id leading a list item ("- REQ-003: ..."), since
    requirements are often bullets. Returns the list-item ids apart so a
    record written before they were counted is compared fairly."""
    h = hashlib.sha256()
    sections, rows, listed = [], {}, set()
    for p in paths:
        data = open(p, "rb").read()
        h.update(p.encode() + b"\0" + data)
        if p.endswith(".md"):
            text = re.sub(
                r"<!--.*?-->", "", data.decode("utf-8", "replace"), flags=re.S
            )
            for line in text.split("\n"):
                m = re.match(r"^#{2,3}\s+(.*)$", line)
                if m:
                    sections.append(m.group(1).strip())
                elif line.startswith("|"):
                    cells = [c.strip() for c in line.strip().strip("|").split("|")]
                    if cells and ID_RE.match(cells[0]):
                        rows[cells[0]] = hashlib.sha256(line.encode()).hexdigest()[:12]
                else:
                    li = LIST_ID_RE.match(line)
                    if li and ID_RE.match(li.group(1)) and li.group(1) not in rows:
                        rows[li.group(1)] = hashlib.sha256(line.encode()).hexdigest()[:12]
                        listed.add(li.group(1))
    return h.hexdigest(), sections, rows, listed


def change_summary(old, sections, rows, listed=()):
    s_old, r_old = set(old.get("sections", [])), old.get("rows", {})
    if not old.get("lists"):  # recorded before list-item ids were counted
        rows = {k: v for k, v in rows.items() if k not in listed}
    added = len(set(sections) - s_old) + len(set(rows) - set(r_old))
    retired = len(s_old - set(sections)) + len(set(r_old) - set(rows))
    expanded = sum(1 for k, v in rows.items() if k in r_old and r_old[k] != v)
    parts = [
        f"{n} {w}"
        for n, w in ((added, "added"), (expanded, "expanded"), (retired, "retired"))
        if n
    ]
    return ", ".join(parts) or "wording changed"


def md_table(path, first_header):
    """The table whose first header cell is first_header, as rows."""
    rows, on = [], False
    for line in open(path, encoding="utf-8"):
        if not line.startswith("|"):
            if on and rows:
                break
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if not on and cells and cells[0].lower() == first_header.lower():
            on = True
            rows.append(cells)
        elif on and not all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
            rows.append(cells)
    return rows


def write_csv(rows, path):
    import csv

    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for r in rows:
            w.writerow([re.sub(r"`", "", c) for c in r])


def git_author():
    try:
        return subprocess.run(
            ["git", "config", "user.name"], capture_output=True, text=True
        ).stdout.strip()
    except OSError:
        return ""


# ------------------------------------------------------------- checklist


def evidence_found(ev):
    for item in ev:
        pat, _, needle = item.partition("|")
        for f in glob.glob(pat, recursive=True):
            if not os.path.isfile(f):
                continue
            if (
                not needle
                or needle in open(f, encoding="utf-8", errors="replace").read()
            ):
                return f
    return ""


def import_checklist(xlsx, state):
    import openpyxl

    ws = openpyxl.load_workbook(xlsx)["Project Checklist"]
    n = 0
    for r in ws.iter_rows(min_row=5, values_only=True):
        if isinstance(r[0], int) and r[2]:
            cur = state.setdefault(r[2], {})
            for k, v in zip(
                ("status", "owner", "target", "remarks"), (r[4], r[5], r[6], r[7])
            ):
                if v not in (None, ""):
                    cur[k] = (
                        str(v)
                        if k != "target"
                        else (v.date().isoformat() if hasattr(v, "date") else str(v))
                    )
            # a status a person chose; an untouched "Not Started" is no choice,
            # and "Completed" is only evidence's when evidence set it
            if r[4] and r[4] != "Not Started" and not (cur.get("by") == "evidence" and r[4] == "Completed"):
                cur["by"] = "person"
            n += 1
    return n


def checklist(out_dir, project, date, state_path, template):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    tpl = read_json(template, None)
    if not tpl or not tpl.get("tasks"):
        raise RuntimeError(f"no tasks in {template}")
    state = read_json(state_path, {})
    auto = 0
    for t in tpl["tasks"]:
        cur = state.setdefault(t["task"], {})
        if cur.get("by") == "person":
            continue
        found = evidence_found(t.get("evidence", []))
        if found:
            cur.update(
                {
                    "status": "Completed",
                    "by": "evidence",
                    "remarks": f"evidence: {found}",
                }
            )
            auto += 1
        elif cur.get("by") == "evidence":
            cur.update({"status": "Not Started", "by": "", "remarks": ""})
    os.makedirs(os.path.dirname(state_path), exist_ok=True)
    json.dump(
        state, open(state_path, "w", encoding="utf-8"), indent=1, ensure_ascii=False
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Checklist"
    ws["A1"] = tpl.get("title", "Project Delivery Checklist")
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = tpl.get("note", "")
    ws.append([])
    head = [
        "S.No",
        "Phase",
        "Task",
        "Responsible",
        "Status",
        "Owner",
        "Target Date",
        "Remarks",
    ]
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F4A7A")
    phase, n = None, 0
    for t in tpl["tasks"]:
        if t["phase"] != phase:
            phase = t["phase"]
            ws.append([phase])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="E8EEF8")
        n += 1
        cur = state.get(t["task"], {})
        ws.append(
            [
                n,
                phase,
                t["task"],
                cur.get("responsible", t["responsible"]),
                cur.get("status", "Not Started"),
                cur.get("owner", ""),
                cur.get("target", ""),
                cur.get("remarks", ""),
            ]
        )
    last = ws.max_row
    statuses = ",".join(tpl.get("statuses", []))
    roles = ",".join(tpl.get("roles", []))
    for col, values in (("E", statuses), ("D", roles)):
        dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        dv.add(f"{col}5:{col}{last}")
        ws.add_data_validation(dv)
    for col, w in zip("ABCDEFGH", (7, 26, 70, 20, 16, 16, 14, 50)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"

    sm = wb.create_sheet("Summary")
    sm["A1"] = "Progress Summary"
    sm["A1"].font = Font(size=14, bold=True)
    sm.append(["Status", "Count"])
    for s in tpl.get("statuses", []):
        sm.append([s, f"=COUNTIF('Project Checklist'!E5:E{last},\"{s}\")"])
    sm.append(["Total Tasks", f"=COUNT('Project Checklist'!A5:A{last})"])
    tot = sm.max_row
    sm.append(
        [
            "% Complete",
            f"=IFERROR(B{3 + tpl['statuses'].index('Completed')}/(B{tot}-B{3 + tpl['statuses'].index('Not Applicable')}),0)",
        ]
    )
    sm.cell(sm.max_row, 2).number_format = "0%"
    sm.append(["Note: % Complete = Completed / (Total - Not Applicable)."])
    sm.append([])
    sm.append(["Phase", "Total", "Completed"])
    for p in dict.fromkeys(t["phase"] for t in tpl["tasks"]):
        sm.append(
            [
                p,
                f"=COUNTIF('Project Checklist'!B5:B{last},\"{p}\")",
                f"=COUNTIFS('Project Checklist'!B5:B{last},\"{p}\",'Project Checklist'!E5:E{last},\"Completed\")",
            ]
        )
    sm.column_dimensions["A"].width = 38
    path = os.path.join(out_dir, f"{project}_DeliveryChecklist_{date}.xlsx")
    wb.save(path)
    return path, n, auto, sum(1 for v in state.values() if v.get("by") == "person")


# ------------------------------------------------------------- the pack


def ignore_pack(out):
    """Keep the generated pack out of git with a rule anchored at the root.

    A bare "deliverables/" line matches at any depth, so it also ignores
    docs/deliverables/, where versions.json and checklist.json must be
    committed. The rule is "/deliverables/" for the default location and the
    pack folder itself for an --out elsewhere (never a whole docs/ tree); a
    bare line for the same folder is rewritten to the anchored form.
    Returns what changed, or "".
    """
    rel = os.path.relpath(out).replace(os.sep, "/")
    top = rel.split("/")[0] if rel.split("/")[0] == "deliverables" else rel
    anchored = f"/{top}/"
    path = ".gitignore"
    lines = open(path, encoding="utf-8").read().splitlines() if os.path.isfile(path) else []
    bare = {top, top + "/"}
    fixed = [anchored if ln.strip() in bare else ln for ln in lines]
    note = ""
    if fixed != lines:
        note = f"rewrote the unanchored {top}/ rule to {anchored}"
    if anchored not in (ln.strip() for ln in fixed):
        fixed.append(anchored)
        note = f"added {anchored}"
    if note:
        first = fixed.index(anchored)
        fixed = [ln for i, ln in enumerate(fixed) if ln != anchored or i == first]
        open(path, "w", encoding="utf-8").write("\n".join(fixed) + "\n")
    return note


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", default="docs/deliverables/pack.json")
    ap.add_argument("--out", default="")
    ap.add_argument("--project", default="")
    ap.add_argument("--customer", default="")
    ap.add_argument("--prepared-by", default="")
    ap.add_argument("--summary", default="")
    ap.add_argument("--date", default=datetime.date.today().isoformat())
    ap.add_argument("--import-checklist", default="")
    ap.add_argument("--no-docx", action="store_true")
    a = ap.parse_args()

    manifest_path = (
        a.manifest
        if os.path.isfile(a.manifest)
        else os.path.join(KIT_TEMPLATES, "pack.json")
    )
    man = read_json(manifest_path, None)
    if not man:
        print(f"pack: cannot read the manifest {manifest_path}", file=sys.stderr)
        return 1
    company = read_json(".bearing/company.json", {})
    project = pascal(a.project or man.get("project") or os.path.basename(os.getcwd()))
    customer = a.customer or man.get("customer", "")
    prepared = (
        a.prepared_by
        or man.get("prepared_by")
        or company.get("delivering_entity")
        or company.get("legal_name", "")
    )
    out = a.out or os.path.join("deliverables", f"{project}_ProjectDocumentation")
    vpath = "docs/deliverables/versions.json"
    versions = read_json(vpath, {"artifacts": {}})

    bad = [
        x["folder"]
        for x in man.get("artifacts", []) + man.get("copy", [])
        if x.get("folder") not in FOLDERS
    ]
    if bad:
        print(
            f"pack: manifest names folders outside the twelve: {', '.join(sorted(set(bad)))}",
            file=sys.stderr,
        )
        return 1

    # a filled checklist is read back before the rebuild, which may delete it
    state_path = "docs/deliverables/checklist.json"
    try:
        if a.import_checklist:
            state = read_json(state_path, {})
            n_imp = import_checklist(a.import_checklist, state)
            os.makedirs(os.path.dirname(state_path), exist_ok=True)
            json.dump(
                state,
                open(state_path, "w", encoding="utf-8"),
                indent=1,
                ensure_ascii=False,
            )
            print(f"pack: imported {n_imp} checklist rows from {a.import_checklist}")
    except Exception as e:
        print(f"pack: cannot import {a.import_checklist}: {e}", file=sys.stderr)
        return 1

    # rebuilt in full: the pack is replaced, never merged; only ever inside
    # this repository and only a folder this script names
    real = os.path.realpath(out)
    if not real.startswith(os.path.realpath(os.getcwd()) + os.sep) or not real.endswith("_ProjectDocumentation"):
        print(f"pack: refusing to replace {out}: the pack lives under this repository and ends in _ProjectDocumentation", file=sys.stderr)
        return 1
    if os.path.isdir(out):
        shutil.rmtree(out)
    for f in FOLDERS:
        os.makedirs(os.path.join(out, f), exist_ok=True)

    written = {f: [] for f in FOLDERS}
    counts = {
        "artifacts": 0,
        "docx": 0,
        "csv": 0,
        "copied": 0,
        "bumped": 0,
        "missing": 0,
    }
    failures = []
    for art in man.get("artifacts", []):
        srcs = expand(art["source"])
        if not srcs:
            counts["missing"] += 1
            print(f"pack: {art['key']}: no source ({art['source']}), left out")
            continue
        counts["artifacts"] += 1
        digest, sections, rows, listed = fingerprint(srcs)
        rec = versions["artifacts"].get(art["key"])
        if rec is None or rec.get("sha256") != digest:
            v = 1 if rec is None else rec["version"] + 1
            summary = a.summary or (
                "First issue" if rec is None else change_summary(rec, sections, rows, listed)
            )
            hist = (rec or {}).get("history", []) + [
                {
                    "version": v,
                    "date": a.date,
                    "author": prepared or git_author(),
                    "summary": summary,
                }
            ]
            rec = {
                "title": art["title"],
                "folder": art["folder"],
                "version": v,
                "sha256": digest,
                "sections": sections,
                "rows": rows,
                "lists": True,
                "history": hist,
            }
            versions["artifacts"][art["key"]] = rec
            counts["bumped"] += 1
        v = rec["version"]
        folder = os.path.join(out, art["folder"])
        for s in srcs:
            dst = os.path.join(folder, os.path.basename(s))
            if s.endswith(".md"):
                # the client's Markdown copy loses HTML comments, as the Word
                # document does: they hold template guidance and internal notes
                text = open(s, encoding="utf-8", errors="replace").read()
                open(dst, "w", encoding="utf-8").write(re.sub(r"<!--.*?-->\n?", "", text, flags=re.S))
            else:
                shutil.copy2(s, dst)
            written[art["folder"]].append(os.path.basename(s))
        if art.get("docx", True) and srcs[0].endswith(".md") and not a.no_docx:
            name = f"{project}_{art['key']}_v{v}.docx"
            try:
                md2docx.build(
                    srcs[0],
                    os.path.join(folder, name),
                    art["title"],
                    project,
                    customer,
                    prepared,
                    v,
                    rec["history"][-1]["date"],
                    art.get("about", ""),
                    list(reversed(rec["history"])),
                )
                counts["docx"] += 1
                written[art["folder"]].append(name)
            except Exception as e:  # a pack with a missing Word document is not a pack
                failures.append(f"{name}: {e}")
        for c in art.get("csv", []):
            rows_ = []
            for s in srcs:
                rows_ += (
                    md_table(s, c["table"])
                    if not rows_
                    else md_table(s, c["table"])[1:]
                )
            if len(rows_) > 1:
                write_csv(rows_, os.path.join(folder, c["file"]))
                counts["csv"] += 1
                written[art["folder"]].append(c["file"])
            else:
                print(
                    f"pack: {art['key']}: no table headed '{c['table']}' for {c['file']}, not written"
                )
        print(f"pack: {art['folder']}/{art['key']} v{v} ({len(srcs)} source files)")

    for cp in man.get("copy", []):
        for s in expand(cp["from"]):
            base = cp.get("strip", "")
            rel = (
                os.path.relpath(s, base)
                if base and s.startswith(base)
                else os.path.basename(s)
            )
            dst = os.path.join(out, cp["folder"], rel)
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            shutil.copy2(s, dst)
            counts["copied"] += 1
            written[cp["folder"]].append(rel)

    if counts["artifacts"] == 0:
        print(
            "pack: 0 artifact sources found in the repository, nothing packed",
            file=sys.stderr,
        )
        return 1

    # folder READMEs
    holds = man.get("folders", {})
    for f in FOLDERS:
        info = holds.get(f, {})
        files = sorted(set(written[f]))
        lines = [
            f"# {f.replace('_', ' ', 1)}",
            "",
            info.get("holds", ""),
            "",
            f"Filled by: {info.get('filled_by', 'the team, by hand')}.",
            "",
        ]
        lines += ["## Files", ""] + ([f"- `{x}`" for x in files] or ["Nothing yet."])
        open(os.path.join(out, f, "README.md"), "w", encoding="utf-8").write(
            "\n".join(lines) + "\n"
        )

    # top README and CHANGELOG
    arts = versions["artifacts"]
    present = {art["key"] for art in man.get("artifacts", []) if expand(art["source"])}
    readme = [
        f"# {project}: Project Documentation",
        "",
        f"{customer or 'Customer not recorded'} · delivered by {prepared or 'not recorded'}",
        "",
        "Generated by client-deliverables from the repository. **Upload the whole folder to the project's Drive.** "
        "It is rebuilt in full each time, so replace rather than merge; CHANGELOG.md carries the history.",
        "",
        "## What is here",
        "",
        "| Folder | Holds | Files |",
        "| --- | --- | --- |",
    ]
    for f in FOLDERS:
        n = len(set(written[f]))
        readme.append(
            f"| `{f}` | {holds.get(f, {}).get('holds', '')} | {n or 'none'} |"
        )
    readme += [
        "",
        "## Current versions",
        "",
        "| Artifact | Version | Last change |",
        "| --- | --- | --- |",
    ]
    for k, r in sorted(arts.items(), key=lambda kv: (kv[1]["folder"], kv[0])):
        if k in present:
            last = r["history"][-1]
            readme.append(
                f"| {r['title']} | v{r['version']} | {last['date']}: {last['summary']} |"
            )
    open(os.path.join(out, "README.md"), "w", encoding="utf-8").write(
        "\n".join(readme) + "\n"
    )
    log = [
        f"# {project}: Changelog",
        "",
        "Every version of every artifact, newest first.",
        "",
    ]
    entries = [
        (h["date"], r["title"], h)
        for k, r in arts.items()
        if k in present
        for h in r["history"]
    ]
    for date, title, h in sorted(
        entries, key=lambda e: (e[0], e[2]["version"]), reverse=True
    ):
        log.append(
            f"- {date}: {title} v{h['version']}, {h['summary']} ({h.get('author') or 'unknown'})"
        )
    open(os.path.join(out, "CHANGELOG.md"), "w", encoding="utf-8").write(
        "\n".join(log) + "\n"
    )
    os.makedirs(os.path.dirname(vpath), exist_ok=True)
    json.dump(
        versions, open(vpath, "w", encoding="utf-8"), indent=1, ensure_ascii=False
    )

    # the delivery checklist
    template = man.get("checklist") or os.path.join(
        KIT_TEMPLATES, "delivery-checklist.json"
    )
    cl = "checklist: not written"
    try:
        path, n, auto, person = checklist(out, project, a.date, state_path, template)
        cl = f"checklist: {n} tasks, {auto} completed from evidence, {person} set by people"
        print(f"pack: {path}")
    except Exception as e:
        failures.append(f"delivery checklist: {e}")

    try:
        note = ignore_pack(out)
        if note:
            print(f"pack: .gitignore: {note}")
    except OSError as e:
        failures.append(f".gitignore: {e}")

    for f in failures:
        print(f"problem: {f}")
    print(
        f"pack: {counts['artifacts']} artifacts ({counts['bumped']} new versions, {counts['missing']} without a source), "
        f"{counts['docx']} Word documents, {counts['csv']} CSV, {counts['copied']} files copied, {cl}; {out}"
    )
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
